"""
document_loader 单元测试：多格式接入与统一页面结构。
"""

import fitz

import pytest

from modules.document_loader import (
    SUPPORTED_EXTENSIONS,
    is_paged_format,
    is_supported,
    load_document,
)

# 生成中文 PDF 必须使用内置 CJK 字体，否则中文会被渲染成乱码
CJK_FONT = "china-s"


def _make_docx(path, paragraphs):
    from docx import Document

    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    document.save(str(path))
    return str(path)


def _make_pptx(path, slides):
    from pptx import Presentation

    presentation = Presentation()
    for lines in slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = lines[0]
        slide.placeholders[1].text = "\n".join(lines[1:])
    presentation.save(str(path))
    return str(path)


def _make_pdf(path, lines):
    doc = fitz.open()
    page = doc.new_page()
    y = 100
    for text, fontsize in lines:
        page.insert_text((72, y), text, fontsize=fontsize, fontname=CJK_FONT)
        y += fontsize + 24
    doc.save(str(path))
    doc.close()
    return str(path)


class TestSupportDetection:
    def test_supported_extensions(self):
        assert ".pdf" in SUPPORTED_EXTENSIONS
        assert ".docx" in SUPPORTED_EXTENSIONS
        assert ".pptx" in SUPPORTED_EXTENSIONS
        assert ".md" in SUPPORTED_EXTENSIONS
        assert ".txt" in SUPPORTED_EXTENSIONS

    def test_is_supported(self):
        assert is_supported("a.pdf") is True
        assert is_supported("A.DOCX") is True
        assert is_supported("a.exe") is False

    def test_is_paged_format(self):
        assert is_paged_format("a.pdf") is True
        assert is_paged_format("a.pptx") is True
        assert is_paged_format("a.docx") is False


class TestLoadDocument:
    def test_load_pdf_delegates(self, tmp_path):
        path = _make_pdf(
            tmp_path / "lecture.pdf",
            [
                ("第一章 检索增强生成", 20),
                ("本章介绍检索增强生成技术在课程资料问答系统中的具体应用方式。", 11),
            ],
        )

        pages = load_document(path)

        assert len(pages) == 1
        assert pages[0]["source"] == "lecture.pdf"
        assert "检索增强生成" in pages[0]["text"]

    def test_load_docx(self, tmp_path):
        path = _make_docx(
            tmp_path / "note.docx",
            ["本章介绍向量数据库的基本原理。", "内容涵盖索引结构与相似度检索方式。"],
        )

        pages = load_document(path)

        assert len(pages) == 1
        assert pages[0]["source"] == "note.docx"
        assert "向量数据库" in pages[0]["text"]
        assert "索引结构" in pages[0]["text"]

    def test_load_markdown(self, tmp_path):
        path = tmp_path / "note.md"
        path.write_text("# 第一章 概述\n\n这是 Markdown 正文内容。", encoding="utf-8")

        pages = load_document(str(path))

        assert len(pages) == 1
        assert pages[0]["source"] == "note.md"
        assert "第一章 概述" in pages[0]["text"]

    def test_load_txt(self, tmp_path):
        path = tmp_path / "note.txt"
        path.write_text("这是纯文本讲义内容。", encoding="utf-8")

        pages = load_document(str(path))

        assert pages[0]["source"] == "note.txt"
        assert "纯文本讲义" in pages[0]["text"]

    def test_load_pptx_maps_slides_to_pages(self, tmp_path):
        path = _make_pptx(
            tmp_path / "slides.pptx",
            [
                ["检索增强生成概述", "用于课程资料问答"],
                ["向量数据库", "加速大规模向量的近似检索过程"],
            ],
        )

        pages = load_document(path)

        assert len(pages) == 2
        # 每张幻灯片对应一个「页」，方便引用溯源
        assert [page["page"] for page in pages] == [1, 2]
        assert pages[0]["source"] == "slides.pptx"
        assert "检索增强生成概述" in pages[0]["text"]
        assert "向量数据库" in pages[1]["text"]


class TestErrorHandling:
    def test_missing_file_raises(self, tmp_path):
        with pytest.raises(FileNotFoundError):
            load_document(str(tmp_path / "missing.pdf"))

    def test_unsupported_extension_raises(self, tmp_path):
        path = tmp_path / "data.bin"
        path.write_bytes(b"\x00\x01")

        with pytest.raises(ValueError):
            load_document(str(path))
