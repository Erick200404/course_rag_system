"""
统一文档接入层。

把不同格式的文件归一化为与 pdf_loader 相同的页面结构：

    {"page": 页码, "text": 文本, "source": 文件名}

选型说明（这里刻意没有使用 langchain-community 的 loader）：
- langchain-community 已进入 sunset，导入会打印弃用警告，不适合放进长期维护的工程
- docx / pptx / md / txt 的解析逻辑本身很轻，用原生库更可控、依赖更少
- 因此 LangChain 的使用集中在真正有收益的 LLM 编排层（见 modules/llm_backends.py 的 LCEL 链）

各格式实现：
- pdf   -> 自研 pymupdf4llm 后端（保留标题层级与表格，效果优于通用 loader）
- docx  -> python-docx（段落 + 表格文本）
- pptx  -> python-pptx，按幻灯片提取，每张幻灯片对应一个「页」
- md/txt-> 直接读取文本
- 无页码概念的格式（docx / md / txt）整份文档记为第 1 页
"""

from __future__ import annotations

from pathlib import Path
from typing import List

from modules.pdf_loader import load_pdf

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".pptx", ".md", ".txt")

# 无分页概念的格式
_UNPAGED_EXTENSIONS = (".docx", ".md", ".txt")


def is_supported(filename: str) -> bool:
    """判断文件扩展名是否被支持。"""
    return Path(filename).suffix.lower() in SUPPORTED_EXTENSIONS


def _make_page(text: str, source: str, page: int = 1) -> dict:
    """按统一结构构造一页数据。"""
    return {
        "page": page,
        "text": text.strip(),
        "source": source,
    }


def _load_docx(path: str, source: str) -> List[dict]:
    """使用 python-docx 读取 Word 文档（段落 + 表格文本）。"""
    from docx import Document

    document = Document(path)

    lines = [paragraph.text.strip() for paragraph in document.paragraphs]

    # 表格内容往往承载关键知识点，一并提取
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(" | ".join(cells))

    text = "\n".join(line for line in lines if line).strip()
    if not text:
        return []

    return [_make_page(text, source)]


def _load_text(path: str, source: str) -> List[dict]:
    """读取 Markdown / 纯文本。"""
    text = Path(path).read_text(encoding="utf-8").strip()
    if not text:
        return []

    return [_make_page(text, source)]


def _load_pptx(path: str, source: str) -> List[dict]:
    """
    使用 python-pptx 按幻灯片提取文本。

    每张幻灯片记为一个「页」，使引用溯源对课件同样有效。
    """
    from pptx import Presentation

    presentation = Presentation(path)

    pages = []
    for index, slide in enumerate(presentation.slides, start=1):
        lines = []

        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue

            for paragraph in shape.text_frame.paragraphs:
                text = "".join(run.text for run in paragraph.runs).strip()
                if text:
                    lines.append(text)

        text = "\n".join(lines).strip()
        if not text:
            continue

        pages.append(_make_page(text, source, page=index))

    return pages


def load_document(path: str) -> List[dict]:
    """
    读取单个文件，按「页」返回统一结构的数据。

    参数:
        path: 文件路径（需在 data 目录或任意本地路径）

    返回:
        [{"page": 页码, "text": 文本, "source": 文件名}, ...]
    """
    file_path = Path(path)

    if not file_path.exists():
        raise FileNotFoundError(f"文件不存在: {path}")

    suffix = file_path.suffix.lower()
    source = file_path.name

    if suffix == ".pdf":
        return load_pdf(path)

    if suffix == ".docx":
        return _load_docx(path, source)

    if suffix in (".md", ".txt"):
        return _load_text(path, source)

    if suffix == ".pptx":
        return _load_pptx(path, source)

    raise ValueError(
        f"不支持的文件格式: {suffix}，当前支持 {', '.join(SUPPORTED_EXTENSIONS)}"
    )


def is_paged_format(filename: str) -> bool:
    """
    判断格式是否天然带页码（用于前端展示引用时的措辞）。
    """
    return Path(filename).suffix.lower() not in _UNPAGED_EXTENSIONS
